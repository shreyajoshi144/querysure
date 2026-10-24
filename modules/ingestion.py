"""Deterministic ingestion/normalization for QuerySure's four supported formats."""
from __future__ import annotations

import csv
import json
import os
import re
import sqlite3
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from config import MAX_INGEST_ROWS, MAX_INGEST_COLUMNS, MAX_INGEST_SHEETS, MAX_XLSX_UNCOMPRESSED_MB, MAX_XLSX_COMPRESSION_RATIO

import pandas as pd
from .data_normalizer import normalize_tabular_values

SUPPORTED_EXTENSIONS = {".csv", ".json", ".xlsx", ".xlsm", ".db", ".sqlite", ".sqlite3"}


def sanitize_filename(filename: str | None) -> str:
    raw = os.path.basename(filename or "upload.db").replace("\x00", "")
    return raw[:255] or "upload.db"


def safe_table_name(name: str, fallback: str = "uploaded_data") -> str:
    base = re.sub(r"[^a-zA-Z0-9_]+", "_", Path(name).stem).strip("_").lower()[:50]
    return base or fallback


def normalize_columns(df: pd.DataFrame, percent_suffix: bool = False) -> pd.DataFrame:
    """Keep every meaningful source column while producing deterministic names.

    Blank/parser-generated headers are renamed rather than dropped so source data
    is never silently discarded. Duplicate headers are suffixed deterministically.
    """
    df = df.copy()
    used: set[str] = set()
    out: list[str] = []
    keep: list[int] = []
    unnamed = re.compile(r"^Unnamed:\s*\d+$", re.I)
    for idx, raw in enumerate(df.columns, start=1):
        name = str(raw).strip()
        if not name or unnamed.match(name):
            # No header AND no data in the whole column: an artifact of a trailing
            # delimiter / empty spreadsheet column, not a field. (A placeholder
            # column that DOES hold values is kept and named column_<n>.)
            if df.iloc[:, idx - 1].isna().all():
                continue
            name = f"column_{idx}"
        if percent_suffix:
            # "Growth %" -> growth_pct instead of silently dropping the unit.
            name = re.sub(r"\s*%", " pct", name)
        name = re.sub(r"\s+", "_", name)
        name = re.sub(r"[^a-zA-Z0-9_]+", "_", name).strip("_").lower() or f"column_{idx}"
        candidate = name
        suffix = 2
        while candidate in used:
            candidate = f"{name}_{suffix}"
            suffix += 1
        used.add(candidate)
        out.append(candidate)
        keep.append(idx - 1)
    if not keep:
        raise ValueError("File appears to have no columns or rows.")
    df = df.iloc[:, keep]
    df.columns = out
    return df


def _sniff_csv(path: str) -> tuple[str, str]:
    encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
    last: Exception | None = None
    for enc in encodings:
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                sample = f.read(128 * 1024)
            if not sample.strip():
                raise ValueError("CSV file is empty.")
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                delimiter = dialect.delimiter
            except csv.Error:
                delimiter = ","
            return enc, delimiter
        except UnicodeDecodeError as exc:
            last = exc
    raise ValueError(f"CSV encoding could not be decoded: {last}")


def read_csv(path: str) -> pd.DataFrame:
    encoding, delimiter = _sniff_csv(path)
    try:
        first = None
        second = None
        row_count = 0
        with open(path, "r", encoding=encoding, newline="") as f:
            for row in csv.reader(f, delimiter=delimiter):
                if not any(str(v).strip() for v in row):
                    continue
                if len(row) > MAX_INGEST_COLUMNS:
                    raise ValueError(f"CSV has more than {MAX_INGEST_COLUMNS} columns.")
                row_count += 1
                if row_count > MAX_INGEST_ROWS:
                    raise ValueError(f"CSV has more than {MAX_INGEST_ROWS} rows.")
                if first is None:
                    first = row
                elif second is None:
                    second = row
                if first is not None and len(row) > len(first) and any(str(c).strip() for c in row[len(first):]):
                    raise ValueError(
                        f"CSV row {row_count} has {len(row)} fields but the header has {len(first)}. "
                        "A value probably contains an unquoted delimiter; quote it or fix the row."
                    )
        if first is None:
            raise ValueError("CSV file is empty.")
        def numeric_like(v: str) -> bool:
            try:
                float(v.strip())
                return True
            except Exception:
                return False
        headerless = bool(second) and all(numeric_like(v) for v in first if str(v).strip()) and all(numeric_like(v) for v in second if str(v).strip())
        if headerless:
            df = pd.read_csv(path, sep=delimiter, encoding=encoding, engine="python", header=None)
            df.columns = [f"column_{i}" for i in range(1, len(df.columns) + 1)]
        else:
            df = pd.read_csv(path, sep=delimiter, encoding=encoding, engine="python", index_col=False)
        numeric_pos = [i for i, c in enumerate(df.columns) if pd.api.types.is_numeric_dtype(df[c])]
        if numeric_pos:
            raw = pd.read_csv(
                path, sep=delimiter, encoding=encoding, engine="python", index_col=False,
                header=None if headerless else 0, dtype=str, keep_default_na=False, usecols=numeric_pos,
            )
            for out_pos, col_pos in enumerate(numeric_pos):
                text = raw.iloc[:, out_pos]
                if text.str.match(r"^\s*[+-]?0\d").any():
                    df[df.columns[col_pos]] = text.where(text != "", None)
    except Exception as exc:
        raise ValueError(f"CSV parsing failed: {exc}") from exc
    if df.shape[1] == 0:
        raise ValueError("CSV contains no columns.")
    return normalize_tabular_values(normalize_columns(df, percent_suffix=True))


def _is_scalar(v: Any) -> bool:
    return v is None or isinstance(v, (str, int, float, bool))


def read_json(path: str) -> pd.DataFrame:
    """Stream JSON records to a bounded temporary JSONL file before DataFrame creation."""
    decoder = json.JSONDecoder()
    staged_path: str | None = None

    def stage_record(record: dict[str, Any], out) -> None:
        if len(record) > MAX_INGEST_COLUMNS:
            raise ValueError(f"JSON record contains more than {MAX_INGEST_COLUMNS} columns.")
        row = {}
        for key, value in record.items():
            if isinstance(value, (dict, list)):
                row[key] = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            elif _is_scalar(value):
                row[key] = value
            else:
                row[key] = str(value)
        out.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")

    row_count = 0
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", suffix=".jsonl", delete=False) as staged:
            staged_path = staged.name
            try:
                with open(path, "r", encoding="utf-8-sig") as f:
                    buffer = f.read(64 * 1024)
                    while buffer and buffer.lstrip() == "":
                        buffer = f.read(64 * 1024)
                    if not buffer:
                        raise ValueError("JSON record array is empty.")
                    buffer = buffer.lstrip()

                    def consume_array(buf: str) -> str:
                        nonlocal row_count
                        while True:
                            buf = buf.lstrip()
                            if not buf:
                                buf = f.read(64 * 1024)
                                if not buf:
                                    raise ValueError("Malformed JSON: unexpected end of document.")
                                continue
                            if buf[0] == "]":
                                return buf[1:]
                            if buf[0] == ",":
                                buf = buf[1:]
                                continue
                            try:
                                record, idx = decoder.raw_decode(buf)
                            except json.JSONDecodeError:
                                more = f.read(64 * 1024)
                                if not more:
                                    raise ValueError("Malformed JSON: unexpected end of document.")
                                buf += more
                                continue
                            if not isinstance(record, dict):
                                raise ValueError("JSON record arrays must contain objects.")
                            if len(record) > MAX_INGEST_COLUMNS:
                                raise ValueError(f"JSON record contains more than {MAX_INGEST_COLUMNS} columns.")
                            row_count += 1
                            if row_count > MAX_INGEST_ROWS:
                                raise ValueError(f"JSON contains more than {MAX_INGEST_ROWS} rows.")
                            stage_record(record, staged)
                            buf = buf[idx:]

                    if buffer[0] == "[":
                        buffer = consume_array(buffer[1:])
                    else:
                        if buffer[0] != "{":
                            raise ValueError("JSON must be an array of records or an object containing one record array.")
                        found_array = False
                        buffer = buffer[1:]
                        while True:
                            buffer = buffer.lstrip()
                            if not buffer:
                                buffer = f.read(64 * 1024)
                                if not buffer:
                                    raise ValueError("Malformed JSON: unexpected end of document.")
                                continue
                            if buffer[0] == "}":
                                buffer = buffer[1:]
                                break
                            if buffer[0] == ",":
                                buffer = buffer[1:]
                                continue
                            try:
                                key, idx = decoder.raw_decode(buffer)
                            except json.JSONDecodeError:
                                more = f.read(64 * 1024)
                                if not more:
                                    raise ValueError("Malformed JSON: unexpected end of document.")
                                buffer += more
                                continue
                            if not isinstance(key, str):
                                raise ValueError("Malformed JSON object.")
                            buffer = buffer[idx:].lstrip()
                            if not buffer or buffer[0] != ":":
                                if not buffer:
                                    buffer += f.read(64 * 1024)
                                    continue
                                raise ValueError("Malformed JSON object.")
                            buffer = buffer[1:].lstrip()
                            if not buffer:
                                buffer = f.read(64 * 1024)
                                continue
                            if buffer[0] == "[":
                                if found_array:
                                    raise ValueError("Ambiguous JSON structure: expected exactly one top-level record array. Multiple collections are rejected to prevent silent data loss.")
                                found_array = True
                                buffer = consume_array(buffer[1:])
                            else:
                                try:
                                    _, idx = decoder.raw_decode(buffer)
                                except json.JSONDecodeError:
                                    more = f.read(64 * 1024)
                                    if not more:
                                        raise ValueError("Malformed JSON: unexpected end of document.")
                                    buffer += more
                                    continue
                                buffer = buffer[idx:]
                        if not found_array:
                            raise ValueError("JSON must be an array of records or an object containing one record array.")

                    if row_count == 0:
                        raise ValueError("JSON record array is empty.")
                    if buffer.strip():
                        raise ValueError("Malformed JSON: trailing data after the record array.")
                    while True:
                        tail = f.read(64 * 1024)
                        if not tail:
                            break
                        if tail.strip():
                            raise ValueError("Malformed JSON: trailing data after the record array.")
            except UnicodeDecodeError as exc:
                raise ValueError("JSON must be UTF-8 encoded.") from exc
            except ValueError:
                raise
            except json.JSONDecodeError as exc:
                raise ValueError(f"Malformed JSON: {exc}") from exc
            except Exception as exc:
                raise ValueError(f"Malformed JSON: {exc}") from exc

        # The row/column limits have already been enforced while streaming, so
        # DataFrame creation happens only after the complete input is validated.
        return normalize_columns(pd.read_json(staged_path, lines=True))
    finally:
        if staged_path:
            try:
                os.unlink(staged_path)
            except FileNotFoundError:
                pass


def _sheet_header_row(values: list[list[Any]]) -> int | None:
    """Find a likely header row while ignoring leading titles/blank rows."""
    nonempty_rows = []
    for idx, row in enumerate(values[:40]):
        cells = [v for v in row if v is not None and str(v).strip() != ""]
        if not cells:
            continue
        unique = {str(v).strip().lower() for v in cells}
        textish = sum(isinstance(v, str) for v in cells)
        next_rows = values[idx + 1:idx + 4]
        has_following_data = any(any(v is not None and str(v).strip() != "" for v in r) for r in next_rows)
        if len(cells) >= 2 and len(unique) == len(cells) and textish >= max(1, len(cells) // 2) and has_following_data:
            return idx
        nonempty_rows.append((idx, cells, textish, has_following_data))
    # Single-column sheets: accept a textual first meaningful row followed by data.
    for idx, cells, textish, has_following_data in nonempty_rows:
        if len(cells) == 1 and textish == 1 and has_following_data:
            return idx
    return None


_FORMULA_SCAN_ROWS = 20000


def _assert_formulas_have_cached_values(path: str, title: str, values: list[list[Any]]) -> None:
    """A workbook written by a library (openpyxl, xlsxwriter, ...) stores formulas
    but no results until a spreadsheet app recalculates and saves it. We read
    cached results, so every formula cell would silently load as empty. Refuse
    that case explicitly. (Only when NO formula in the sheet has a cached value:
    a formula that legitimately evaluates to an empty string must not trip it.)"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[title]
        formulas = uncached = 0
        for r_idx, (f_row, v_row) in enumerate(zip(ws.iter_rows(values_only=True), values)):
            if r_idx >= _FORMULA_SCAN_ROWS:
                break
            for c_idx, f in enumerate(f_row):
                if isinstance(f, str) and f.startswith("="):
                    formulas += 1
                    if c_idx >= len(v_row) or v_row[c_idx] is None:
                        uncached += 1
        if formulas and uncached == formulas:
            raise ValueError(
                f"Sheet '{title}' contains {formulas} formula cell(s) with no saved results "
                "(the file was generated by a program and never recalculated). Open it in "
                "Excel or LibreOffice, save it, and upload again."
            )
    finally:
        wb.close()


def _guard_xlsx_container(path: str) -> None:
    try:
        with zipfile.ZipFile(path, "r") as zf:
            infos = zf.infolist()
            total_uncompressed = sum(max(0, i.file_size) for i in infos)
            limit = MAX_XLSX_UNCOMPRESSED_MB * 1024 * 1024
            if total_uncompressed > limit:
                raise ValueError(f"Excel workbook expands beyond the {MAX_XLSX_UNCOMPRESSED_MB} MB safety limit.")
            for info in infos:
                if info.file_size > 0 and info.compress_size > 0:
                    ratio = info.file_size / info.compress_size
                    if ratio > MAX_XLSX_COMPRESSION_RATIO:
                        raise ValueError("Excel workbook compression ratio is unsafe and may indicate a zip bomb.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Excel workbook is not a valid ZIP container.") from exc


def read_excel(path: str) -> list[tuple[str, pd.DataFrame]]:
    _guard_xlsx_container(path)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=path.lower().endswith(".xlsm"))
    except Exception as exc:
        raise ValueError(f"Excel workbook is malformed or unreadable: {exc}") from exc

    tables: list[tuple[str, pd.DataFrame]] = []
    if len(wb.worksheets) > MAX_INGEST_SHEETS:
        wb.close()
        raise ValueError(f"Excel workbook contains more than {MAX_INGEST_SHEETS} sheets.")
    for ws in wb.worksheets:
        values = []
        row_count = 0
        max_columns = 0
        for row in ws.iter_rows(values_only=True):
            row_count += 1
            max_columns = max(max_columns, len(row))
            if max_columns > MAX_INGEST_COLUMNS:
                wb.close()
                raise ValueError(f"Sheet '{ws.title}' has more than {MAX_INGEST_COLUMNS} columns.")
            if row_count > MAX_INGEST_ROWS:
                wb.close()
                raise ValueError(f"Sheet '{ws.title}' has more than {MAX_INGEST_ROWS} rows.")
            if len(values) < 40:
                values.append(list(row))
        if not any(any(v is not None and str(v).strip() != "" for v in row) for row in values):
            continue
        header = _sheet_header_row(values)
        if header is None:
            raise ValueError(f"Sheet '{ws.title}' has data but no unambiguous header row.")
        _assert_formulas_have_cached_values(path, ws.title, values)
        try:
            df = pd.read_excel(path, sheet_name=ws.title, header=header, engine="openpyxl")
            # pandas turns the TEXT cell "00123" into 123. Re-read as text only to
            # find numeric columns whose source cells had leading zeros, and keep
            # those as text (zip codes, account / phone / ID numbers).
            numeric_cols = [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])]
            if numeric_cols:
                raw = pd.read_excel(path, sheet_name=ws.title, header=header, engine="openpyxl", dtype=str)
                for col in numeric_cols:
                    if col in raw.columns:
                        text = raw[col]
                        if text.fillna("").str.match(r"^\s*[+-]?0\d").any():
                            df[col] = text
        except Exception as exc:
            raise ValueError(f"Could not parse Excel sheet '{ws.title}': {exc}") from exc
        if df.shape[1] == 0:
            continue
        df = normalize_tabular_values(normalize_columns(df, percent_suffix=True))
        tables.append((ws.title, df))
    wb.close()
    if not tables:
        raise ValueError("Excel workbook contains no usable non-empty sheets.")
    return tables


def validate_sqlite(path: str) -> list[str]:
    try:
        conn = sqlite3.connect(f"file:{Path(path).resolve()}?mode=ro", uri=True)
        try:
            integrity = conn.execute("PRAGMA integrity_check;").fetchone()[0]
            if str(integrity).lower() != "ok":
                raise ValueError(f"SQLite integrity check failed: {integrity}")
            tables = [r[0] for r in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
            ).fetchall()]
            if not tables:
                raise ValueError("SQLite database contains no usable tables.")
            if len(tables) > MAX_INGEST_SHEETS:
                raise ValueError(f"SQLite database contains more than {MAX_INGEST_SHEETS} tables.")
            for table in tables:
                columns = conn.execute(f'SELECT COUNT(*) FROM pragma_table_info("{table.replace(chr(34), chr(34)*2)}")').fetchone()[0]
                if columns > MAX_INGEST_COLUMNS:
                    raise ValueError(f"SQLite table '{table}' has more than {MAX_INGEST_COLUMNS} columns.")
                rows = conn.execute(f'SELECT COUNT(*) FROM "{table.replace(chr(34), chr(34)*2)}"').fetchone()[0]
                if rows > MAX_INGEST_ROWS:
                    raise ValueError(f"SQLite table '{table}' has more than {MAX_INGEST_ROWS} rows.")
            return tables
        finally:
            conn.close()
    except sqlite3.DatabaseError as exc:
        raise ValueError(f"File is not a valid SQLite database: {exc}") from exc


def ingest_to_sqlite(source_path: str, extension: str, destination_db: str, base_name: str) -> list[str]:
    """Normalize one supported source into a SQLite dataset; returns table names."""
    ext = extension.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported format: {ext}")
    if ext in {".db", ".sqlite", ".sqlite3"}:
        validate_sqlite(source_path)
        return validate_sqlite(source_path)

    conn = sqlite3.connect(destination_db)
    tables: list[str] = []
    try:
        if ext == ".csv":
            df = read_csv(source_path)
            name = safe_table_name(base_name)
            df.to_sql(name, conn, if_exists="replace", index=False)
            tables.append(name)
        elif ext == ".json":
            df = read_json(source_path)
            name = safe_table_name(base_name)
            df.to_sql(name, conn, if_exists="replace", index=False)
            tables.append(name)
        else:
            excel_tables = read_excel(source_path)
            base = safe_table_name(base_name)
            for sheet_name, df in excel_tables:
                if len(excel_tables) == 1:
                    name = base
                else:
                    sheet = safe_table_name(sheet_name, fallback="sheet")
                    # Reserve room for collision suffixes so truncation can never
                    # make two sheet names identical or create an infinite loop.
                    base_part = base[:30]
                    sheet_part = sheet[:20]
                    name = f"{base_part}_{sheet_part}"[:60]
                    suffix = 2
                    while name in tables:
                        name = f"{base_part}_{sheet_part}_{suffix}"[:60]
                        suffix += 1
                df.to_sql(name, conn, if_exists="replace", index=False)
                tables.append(name)
        conn.commit()
    finally:
        conn.close()
    return tables
