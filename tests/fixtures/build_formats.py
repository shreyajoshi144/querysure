"""Deterministic stand-in fixtures (Excel / CSV / JSON) used when the real user
files are not present in tests/fixtures/user/. The ground truth is returned as a
plain-python list so tests can compute expectations independently of the pipeline."""
import json, random
from datetime import date, datetime, timedelta

REGIONS = ["North", "South", "East", "West"]
CATEGORIES = ["Electronics", "Furniture", "Office Supplies"]


def make_rows(n=120, seed=7):
    rnd = random.Random(seed)
    rows = []
    start = date(2023, 1, 1)
    for i in range(n):
        d = start + timedelta(days=rnd.randint(0, 364))
        rows.append({
            "order_id": 1000 + i,
            "order_date": d,
            "region": rnd.choice(REGIONS),
            "category": rnd.choice(CATEGORIES),
            "revenue": round(rnd.uniform(100, 5000), 2),
            "units": rnd.randint(1, 40),
            "margin": round(rnd.uniform(0.05, 0.45), 4),       # fraction 0-1
            "growth_pct": round(rnd.uniform(-10, 40), 1),       # percentage points
        })
    # guarantee boundary dates used by range tests
    rows[0]["order_date"] = date(2023, 1, 31)
    rows[1]["order_date"] = date(2023, 1, 1)
    return rows


def write_xlsx(path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["order_id", "order_date", "region", "category", "revenue", "units", "margin", "growth_pct"])
    for r in rows:
        ws.append([r["order_id"], datetime(r["order_date"].year, r["order_date"].month, r["order_date"].day),
                   r["region"], r["category"], r["revenue"], r["units"], r["margin"], r["growth_pct"]])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd"
        row[6].number_format = "0.0%"
    ws2 = wb.create_sheet("Targets")
    ws2.append(["region", "target_revenue"])
    for reg in REGIONS:
        ws2.append([reg, 50000])
    wb.save(path)


def write_csv(path, rows):
    import csv
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Order ID", "Order Date", "Region", "Category", "Revenue", "Units", "Margin %", "Growth %"])
        for r in rows:
            w.writerow([r["order_id"], r["order_date"].strftime("%d/%m/%Y"), r["region"], r["category"],
                        f"${r['revenue']:,.2f}", r["units"], f"{r['margin'] * 100:.2f}%", f"{r['growth_pct']}%"])


def write_json(path, rows):
    data = [{**r, "order_date": r["order_date"].isoformat()} for r in rows]
    with open(path, "w") as f:
        json.dump(data, f)
